"""Rebuild the editable handout using python-docx (instructor utility only)."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOC = Document()
sec = DOC.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.75)
sec.header_distance = sec.footer_distance = Inches(0.3)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = DOC.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
# The bundled starter template may carry a title paragraph border. Remove
# paragraph borders from inherited styles so the title uses typography only.
for style in DOC.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = DOC.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.widow_control = True
DOC.styles['Title'].font.size = Pt(23)
DOC.styles['Title'].font.bold = True
DOC.styles['Heading 1'].font.size = Pt(16)
DOC.styles['Heading 2'].font.size = Pt(11.5)
for name in ['Heading 1','Heading 2']:
    DOC.styles[name].font.bold = True
    DOC.styles[name].paragraph_format.keep_with_next = True
code_style = DOC.styles.add_style('Lab Code', 1)
code_style.font.name = 'Consolas'; code_style.font.size = Pt(9)
code_style.paragraph_format.line_spacing = 1.0
code_style.paragraph_format.space_after = Pt(6)
code_style.paragraph_format.left_indent = Inches(0.12)
code_style.paragraph_format.keep_together = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('DASE4136 MATLAB Lab   |   ')
run.font.name = 'Arial'; run.font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
DOC.core_properties.title = 'DASE4136 Robot Motion Control Using a Lidar Map'
DOC.core_properties.subject = 'Student steps experiments and required evidence'
DOC.core_properties.author = 'DASE4136'


def p(text, bold=False):
    para = DOC.add_paragraph()
    r = para.add_run(text); r.bold = bold
    return para


def h(text, level=2):
    DOC.add_heading(text, level)


def code(text):
    DOC.add_paragraph(text, 'Lab Code')


def step(number, text):
    para = DOC.add_paragraph()
    para.add_run(f'{number}. ').bold = True
    para.add_run(text)
    return para


def table(headers, rows, widths):
    t = DOC.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width); cell.text = text
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, text, width in zip(cells,row,widths):
            c.width = Inches(width); c.text = str(text)
    for i,row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b = OxmlElement('w:'+side)
                for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    b.set(qn('w:'+key),value)
                borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m = OxmlElement('w:'+side); m.set(qn('w:w'),'95'); m.set(qn('w:type'),'dxa'); margins.append(m)
            tcpr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DDE7F0' if i==0 else ('F5F7F9' if i%2==0 else 'FFFFFF'))
            tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for r in para.runs:
                    r.font.name='Arial'; r.font.size=Pt(9.5); r.bold=(i==0)
    DOC.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def page(title):
    DOC.add_page_break(); h(title,1)


def link(label,url):
    para=DOC.add_paragraph()
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hl=OxmlElement('w:hyperlink'); hl.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'145C92'); pr.append(color)
    r.append(pr); txt=OxmlElement('w:t'); txt.text=label; r.append(txt); hl.append(r); para._p.append(hl)


DOC.add_heading('DASE4136 Robot Motion Control Using a Lidar Map',0)
p('Student lab sheet   |   Suggested time 90 minutes',True)
p('Build a map from lidar scans and control a differential-drive robot through that mapped environment. A short RRT planning step provides the reference route. Your main task is to understand the feedback loop and compare two controller settings on the same route.')
h('Software and download')
p('Use MATLAB R2023b or newer with Navigation Toolbox and Robotics System Toolbox installed and licensed. Apple-silicon Macs use native MATLAB R2023b or newer; Intel Macs use a supported Intel release from R2023b through R2025b. Your macOS version must be supported by that release. No ROS, Simulink, compiler or robot hardware is needed.')
step(1,'In Terminal on Mac, or PowerShell/Git Bash on Windows, run the following commands. Git must already be installed.')
code('git clone -b alternative_lab https://github.com/SAS-HKU/DASE4136_student.git\ncd DASE4136_student/alternative_lab')
link('Without Git download the alternative_lab branch ZIP','https://github.com/SAS-HKU/DASE4136_student/archive/refs/heads/alternative_lab.zip')
p('If using the ZIP, extract it and locate the alternative_lab subfolder. The Word sheet and MATLAB files are inside that folder. Use this branch for the MATLAB lab; the course repository also contains other lab materials.')
step(2,'In MATLAB, set Current Folder to the alternative_lab subfolder containing lab_config.m. Run the commands below. Both toolboxes should show Installed = true and LicenseAvailable = true.')
code('setup_lab;\ncheck_environment;\ncfg = lab_config();')
h('The workflow')
table(['Stage','What you will do','Time'],[
 ['Setup','Download and check MATLAB','10 min'],
 ['Map and route','Inspect a lidar map and one feasible route','20 min'],
 ['Motion control','Follow the route and inspect wheel commands','30 min'],
 ['Comparison','Change lookahead and explain the results','30 min']
],[1.3,4.55,0.9])
p('The supplied code runs without TODO implementation blocks. Maps and scans are generated locally, so runtime internet access is unnecessary. Each run saves figures and CSV measurements in output. Complete the steps on pages 2 to 4 and submit the short evidence set listed on page 4.')

page('Build the map and obtain a reference route')
h('Step 1 Inspect the mapping result')
code('mapping = run_slam(cfg);')
p('Open output/baseline/slam.png. The left panel shows the occupancy map and estimated survey trajectory; the right panel shows the pose graph. Identify the room walls and three obstacles. Describe the difference between free, occupied and unknown cells.')
p('Expected: about 129 scans are accepted and a closed survey loop appears. Black cells represent occupied space, white cells free space, and gray cells unknown space. Some obstacle interiors remain unseen. Exact loop scores can differ by MATLAB release.')
p('The map origin is the first lidar scan pose [0 0 0]. All route and robot coordinates below use this same frame. The survey data is simulated; the true survey poses are used only to generate scans and assess mapping error, not as inputs to SLAM.')
h('Step 2 Generate one supporting route')
code('route = run_planning(mapping,cfg);')
p('Open output/baseline/route.png. The start is [0 0 0] and the goal is [10 6 pi/2], in metres and radians. RRT connects vehicle poses using forward-only Dubins curves. Treat this as a reference route for control; you do not need to tune the planner in this lab.')
p('Expected: Success and MotionValid are true. The path lies in observed free space. Unknown cells are blocked, and obstacles are inflated by the robot radius 0.20 m plus a planning margin 0.15 m. Inflation provides room for the robot and for small tracking deviations.')
h('Record two observations')
step(1,'Identify one region of unknown space in the map. Explain why it should not be silently treated as safe to drive through.')
step(2,'Explain why a collision-free planned route does not guarantee that a feedback-controlled robot will follow it exactly.')
h('Files worth opening')
p('run_slam.m builds the map. run_planning.m uses that estimated map rather than a separate perfect map. The planner helper checks each curved connection with isMotionValid. You may inspect these files, but the main code to study is run_control.m on the next page.')
p('If no route is found, restore lab_config and inspect the map and endpoints before changing the iteration budget. Record the failure rather than removing unknown-space or collision checks.')

page('Control the differential drive robot')
h('Step 3 Run the feedback controller')
code('control = run_control(route,cfg);')
p('Watch the robot footprint and heading move through the mapped room. Open control_trajectory.png and control_signals.png in output/baseline. The trajectory plot compares the planned route with the actual robot path; the signal plots show tracking error, requested/actual speed, and left/right wheel commands.')
h('Follow the loop in run_control')
p('At each sample, the current robot pose is fed to a pure pursuit controller. It chooses a point ahead on the route and returns speed v and curvature kappa. The code converts curvature into heading rate omega, then into wheel angular speeds, clips those speeds to the wheel limit and advances the differential-drive model.')
code('omega = v * kappa\nwheelL = (v - omega*b/2) / r\nwheelR = (v + omega*b/2) / r')
p('Here r is wheel radius and b is track width. Speeds v, omega and wheelL/wheelR use m/s, rad/s and rad/s respectively; curvature uses 1/m. The baseline wheel radius is 0.10 m and track width is 0.40 m.')
step(1,'For v = 0.40 m/s and omega = 0.80 rad/s, calculate both wheel commands. Compare them with the +/-8 rad/s wheel limit. Explain which wheel must turn faster for a positive heading rate.')
step(2,'Locate the feedback, wheel-command conversion, robot-model update and stopping condition in run_control.m. Explain what changing cfg.control.lookahead means physically.')
h('Expected baseline result')
table(['Measurement','Expected behavior'],[
 ['Goal status','ReachedGoal = true and CollisionDetected = false'],
 ['Final goal distance','At most 0.15 m'],
 ['Tracking RMSE','About 0.013 m on the reference Windows run'],
 ['Simulated travel time','About 48 s at 0.40 m/s']
],[2.1,4.65])
p('The controller stops using position tolerance; final heading is not regulated. Tracking error is distance to the nearest segment of the reference route. Simulated time is not the time MATLAB takes to compute the run.')
p('This exercise uses the true simulated pose as feedback after mapping is complete. It does not estimate the driving pose online. A collision guard checks the actual robot footprint and stops before accepting a blocked motion; it does not steer around obstacles.')

page('Compare two settings and submit the results')
h('Step 4 Change lookahead on the same route')
p('Before running, predict how a smaller and a larger lookahead distance will affect steering, tracking error and the travelled path. Then run the supplied comparison. It reuses your existing map and route and changes only lookahead.')
code('summary = run_experiments(route,cfg);')
p('The two settings are 0.25 m and 0.80 m. Each saves its control plots and CSV files under output/experiments, and comparison.csv contains both result rows. Read each StopReason; a collision or time-limit stop must not be reported as a successful arrival.')
table(['Lookahead','Reached goal','Tracking RMSE','Maximum error','Simulated time'],[
 ['0.25 m','Record result','Record in m','Record in m','Record in s'],
 ['0.80 m','Record result','Record in m','Record in m','Record in s']
],[1.1,1.25,1.5,1.5,1.4])
step(1,'Compare the trajectory and wheel-command plots. Explain the tradeoff you observe between tracking the route closely and making smoother turns.')
step(2,'State which setting you would choose for this mapped room, using your measurements. Explain why a large lookahead can cut corners even when the reference path itself is valid.')
h('Deliver a small evidence set')
p('Submit a 1 to 2 page report containing your wheel-speed calculation, the comparison table, and your answers to the four explanation prompts on pages 2 and 4. Add the baseline map/route figure, both control plots for each lookahead trial, comparison.csv and your experiment .m script. Combine them in one ZIP file. Record MATLAB release, platform and seed 4136. A failure is a result to explain, not a reason to discard a trial.')
h('Optional extra trial')
code("cfg.control.speed = 0.6;\ncfg.outputDir = fullfile(setup_lab(),'output','speed_test');\ntrial = run_control(route,cfg);")
p('Reuse the same route and explain any change in wheel-speed limits, tracking error or stopping status. The extra trial is optional. To run the complete baseline automatically, use results = run_all; the individual commands above show each stage separately.')
h('Help and verification')
p('Function not found: set Current Folder to the alternative_lab folder and run setup_lab. Missing toolbox: run check_environment and resolve the licensed installation. To verify the complete baseline and collision guard, run validate_lab. Exact routes and errors may vary across supported releases.')
link('Course materials on the alternative_lab branch','https://github.com/SAS-HKU/DASE4136_student/tree/alternative_lab')
link('MathWorks pure pursuit controller reference','https://www.mathworks.com/help/robotics/ref/controllerpurepursuit-system-object.html')
p('Lecture links: D5 pages 2 to 7 for kinematics, D7 pages 2 to 6 for observations, and D8 pages 7 to 8 and 22 for motion planning, control and vehicle constraints. Separate slide assignments are outside this lab.')

destination = ROOT/'docs'/'DASE4136_MATLAB_Lab_Sheet.docx'
destination.parent.mkdir(exist_ok=True)
DOC.save(destination)
print(destination)
